#!/usr/bin/env python3
"""Remove sample-workbook rule definitions already supplied by diagen.defaults.

The originals are retained in ``samples/_backups``. A backup is created only
when absent, so rerunning this command never overwrites the pristine copy.

This utility intentionally uses the default *key* as the duplicate criterion:
the web app has historically preloaded the static palette, so rows with a
built-in entity/relation key were already inactive there. After Phase 1,
workbook rows can override defaults; stripping these known sample rows keeps
the web app's prior rendered output while making its sample workbooks smaller.
"""
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from openpyxl import load_workbook

from diagen.defaults import DEFAULT_LINES, DEFAULT_SHAPES


FULLY_STATIC = {
    'S1_integration_landscape.xlsx',
    'S2_robot_wiring.xlsx',
    'S3_aircraft_avionics.xlsx',
    'S4_global_network.xlsx',
    'S5_order_to_cash.xlsx',
    'S6_data_platform.xlsx',
    'ZZ_multi_diagram_test.xlsx',
}
PARTIAL_STATIC = {'TAKSY_VPM_Landscape_v3.xlsx'}


def nonempty_keys(ws, key: str) -> list[str]:
    header = [str(value).strip() if value is not None else '' for value in next(ws.iter_rows(values_only=True))]
    try:
        index = header.index(key)
    except ValueError as exc:
        raise ValueError(f'{ws.title} is missing its {key!r} column') from exc
    return [str(row[index]).strip() for row in ws.iter_rows(min_row=2, values_only=True)
            if len(row) > index and row[index] not in (None, '')]


def remove_builtin_rows(ws, key: str, builtin: dict) -> int:
    header = [str(cell.value).strip() if cell.value is not None else '' for cell in ws[1]]
    try:
        index = header.index(key) + 1
    except ValueError as exc:
        raise ValueError(f'{ws.title} is missing its {key!r} column') from exc
    rows = [row for row in range(ws.max_row, 1, -1)
            if str(ws.cell(row, index).value or '').strip() in builtin]
    for row in rows:
        ws.delete_rows(row, 1)
    return len(rows)


def backup_samples(samples_dir: Path, dry_run: bool) -> None:
    backup_dir = samples_dir / '_backups'
    if not dry_run:
        backup_dir.mkdir(exist_ok=True)
    for path in sorted(samples_dir.glob('*.xlsx')):
        if path.name.startswith('~$'):
            continue
        backup = backup_dir / path.name
        if backup.exists():
            print(f'backup retained: {backup.name}')
        elif dry_run:
            print(f'would back up: {path.name}')
        else:
            shutil.copy2(path, backup)
            print(f'backed up: {path.name}')


def strip_workbook(path: Path, dry_run: bool) -> list[str]:
    workbook = load_workbook(path)
    changed: list[str] = []

    if path.name in FULLY_STATIC:
        for sheet, key, builtin in (
            ('Shape_Library', 'entity_type', DEFAULT_SHAPES),
            ('Line_Rules', 'relation_type', DEFAULT_LINES),
        ):
            if sheet not in workbook.sheetnames:
                continue
            names = nonempty_keys(workbook[sheet], key)
            if not all(name in builtin for name in names):
                raise ValueError(f'{path.name}: {sheet} contains a custom definition')
            workbook.remove(workbook[sheet])
            changed.append(f'removed {sheet} ({len(names)} built-in rows)')

    elif path.name in PARTIAL_STATIC:
        if 'Shape_Library' in workbook.sheetnames:
            removed = remove_builtin_rows(workbook['Shape_Library'], 'entity_type', DEFAULT_SHAPES)
            if removed:
                changed.append(f'removed {removed} built-in Shape_Library rows')
        if 'Line_Rules' in workbook.sheetnames:
            names = nonempty_keys(workbook['Line_Rules'], 'relation_type')
            if not all(name in DEFAULT_LINES for name in names):
                raise ValueError(f'{path.name}: Line_Rules contains a custom definition')
            workbook.remove(workbook['Line_Rules'])
            changed.append(f'removed Line_Rules ({len(names)} built-in rows)')

    if changed:
        if dry_run:
            print(f'would update {path.name}: ' + '; '.join(changed))
        else:
            workbook.save(path)
            print(f'updated {path.name}: ' + '; '.join(changed))
    else:
        print(f'unchanged: {path.name}')
    return changed


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--samples-dir', type=Path,
                        default=Path(__file__).resolve().parents[1] / 'samples')
    parser.add_argument('--dry-run', action='store_true', help='report changes without writing files')
    args = parser.parse_args()

    backup_samples(args.samples_dir, args.dry_run)
    for name in sorted(FULLY_STATIC | PARTIAL_STATIC):
        path = args.samples_dir / name
        if not path.exists():
            raise FileNotFoundError(path)
        strip_workbook(path, args.dry_run)


if __name__ == '__main__':
    main()
