"""Tests for spec.py: loading rule sheets + Nodes/Edges from workbook or CSV."""
from __future__ import annotations
import pytest
from diagen.spec import Spec


def test_spec_loads_shape_and_line_rules(spec):
    assert len(spec.shapes) > 0
    assert len(spec.lines) > 0
    # every shape/line def must have resolved to real dataclass fields, not blanks
    for s in spec.shapes.values():
        assert s.entity_type
        assert s.fill.startswith('#')


def test_spec_loads_sample_diagrams(spec):
    assert set(spec.diagrams) >= {'D1', 'D2', 'D3', 'D4', 'D5'}
    d2 = spec.diagrams['D2']
    assert d2.nodes and d2.edges


def test_spec_links_parent_child(spec):
    # D1 is the layered EA view - some containers should have children
    d1 = spec.diagrams['D1']
    has_children = any(n.children for n in d1.nodes.values())
    assert has_children


def test_shape_of_raises_on_unknown_entity_type(spec):
    d2 = spec.diagrams['D2']
    n = next(iter(d2.nodes.values()))
    n.type = 'totally_not_a_real_entity_type'
    with pytest.raises(KeyError):
        spec.shape_of(n)


def test_line_of_raises_on_unknown_relation_type(spec):
    d2 = spec.diagrams['D2']
    e = d2.edges[0]
    e.relation = 'totally_not_a_real_relation'
    with pytest.raises(KeyError):
        spec.line_of(e)


def test_load_csv_replaces_workbook_sample_data(spec, tmp_path):
    nodes_csv = tmp_path / 'nodes.csv'
    edges_csv = tmp_path / 'edges.csv'
    nodes_csv.write_text(
        'diagram_id,node_id,parent_id,entity_type,label\n'
        'D1,A,,business_actor,Alpha\n'
        'D1,B,,business_actor,Beta\n', encoding='utf-8')
    edges_csv.write_text(
        'diagram_id,edge_id,source_id,target_id,relation_type\n'
        'D1,e1,A,B,association\n', encoding='utf-8')
    spec.load_csv(str(nodes_csv), str(edges_csv), title='My Diagram')
    assert set(spec.diagrams) == {'D1'}
    assert set(spec.diagrams['D1'].nodes) == {'A', 'B'}
    assert spec.diagrams['D1'].title == 'My Diagram'


def test_load_csv_without_edges_leaves_edges_empty(spec, tmp_path):
    nodes_csv = tmp_path / 'nodes_only.csv'
    nodes_csv.write_text(
        'diagram_id,node_id,parent_id,entity_type,label\n'
        'D1,A,,business_actor,Alpha\n', encoding='utf-8')
    spec.load_csv(str(nodes_csv))
    assert spec.diagrams['D1'].edges == []


def test_node_meta_dict_parses_key_values():
    from diagen.model import Node
    n = Node(id='n', diagram='D1', parent='', type='t', label='L',
              meta='lifecycle=phase-out,owner=platform')
    assert n.meta_dict() == {'lifecycle': 'phase-out', 'owner': 'platform'}


def test_node_rows_parses_erd_rows():
    from diagen.model import Node
    n = Node(id='n', diagram='D1', parent='', type='t', label='Orders',
              meta='rows=id(PK);customer_id;total')
    assert n.rows() == ['id(PK)', 'customer_id', 'total']


def test_node_rows_parses_ports():
    from diagen.model import Node
    n = Node(id='n', diagram='D1', parent='', type='t', label='Battery',
              meta='ports=+24V;GND;CAN_H')
    assert n.rows() == ['+24V', 'GND', 'CAN_H']


def test_template_workbook_is_self_contained(template_spec):
    """Diagram_Authoring_Template.xlsx should render without needing the full
    spec workbook - it ships its own Shape_Library/Line_Rules/Nodes/Edges."""
    assert template_spec.shapes
    assert template_spec.lines
    assert any(d.nodes for d in template_spec.diagrams.values())


def test_workbook_uses_defaults_and_overrides_them_locally(tmp_path):
    """Rule sheets are optional, and any supplied row wins over a default."""
    from openpyxl import Workbook

    path = tmp_path / 'local-rules.xlsx'
    wb = Workbook()
    wb.remove(wb.active)

    shapes = wb.create_sheet('Shape_Library')
    shapes.append(['entity_type', 'family', 'shape', 'default_w', 'default_h',
                   'min_w', 'min_h', 'fill_hex', 'stroke_hex', 'auto_size'])
    shapes.append(['application_component', 'Local', 'rounded rectangle', 180, 70,
                   120, 50, '#123456', '#654321', 'N'])

    lines = wb.create_sheet('Line_Rules')
    lines.append(['relation_type', 'family', 'line_style', 'width_px', 'source_end',
                  'target_end', 'routing', 'label_position', 'color_hex'])
    lines.append(['association', 'Local', 'dashed', 3, 'none', 'filled arrow',
                  'straight', 'near source', '#112233'])

    nodes = wb.create_sheet('Nodes')
    nodes.append(['diagram_id', 'node_id', 'parent_id', 'entity_type', 'label'])
    nodes.append(['D1', 'A', '', 'business_actor', 'Built-in fallback'])
    nodes.append(['D1', 'B', '', 'application_component', 'Workbook override'])

    edges = wb.create_sheet('Edges')
    edges.append(['diagram_id', 'edge_id', 'source_id', 'target_id', 'relation_type'])
    edges.append(['D1', 'E1', 'A', 'B', 'association'])
    wb.save(path)

    loaded = Spec(path)
    assert loaded.shapes['business_actor'].entity_type == 'business_actor'
    assert loaded.shapes['application_component'].fill == '#123456'
    assert loaded.shapes['application_component'].default_w == 180
    assert loaded.lines['association'].style == 'dashed'
    assert loaded.lines['association'].target_end == 'filled arrow'
